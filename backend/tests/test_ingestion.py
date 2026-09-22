"""Every document type the composer and the uploader accept is readable."""

import io

import pytest


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost"])
    ws.append(["Rent", 42000])
    ws.append([None, None])
    wb.create_sheet("Notes").append(["Pay rent on the 1st"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Q3 plan"
    slide.placeholders[1].text = "Ship the composer"
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def _docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Hello from Word.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "a"
    table.rows[0].cells[1].text = "b"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class TestExtractPages:
    def test_a_spreadsheet_is_one_entry_per_sheet_with_cells_kept_together(self):
        from app.services.document_ingestion import extract_pages

        pages = extract_pages(_xlsx(), "xlsx")
        assert [p for p, _ in pages] == [None, None]
        assert pages[0][1] == "Sheet: Budget\nItem | Cost\nRent | 42000"
        assert pages[1][1] == "Sheet: Notes\nPay rent on the 1st"

    def test_a_deck_is_one_entry_per_slide(self):
        from app.services.document_ingestion import extract_pages

        pages = extract_pages(_pptx(), "pptx")
        assert pages[0][0] == 1
        assert "Q3 plan" in pages[0][1] and "Ship the composer" in pages[0][1]

    def test_word_tables_come_through(self):
        from app.services.document_ingestion import extract_pages

        text = extract_pages(_docx(), "docx")[0][1]
        assert "Hello from Word." in text and "a | b" in text

    @pytest.mark.parametrize(
        ("file_type", "raw", "expected"),
        [
            ("csv", b"name,qty\r\nbolt,3\r\n", "name | qty\nbolt | 3"),
            ("tsv", b"name\tqty\nbolt\t3\n", "name | qty\nbolt | 3"),
            ("html", b"<html><style>x{}</style><body><h1>Hi</h1><p>There</p></body></html>", "Hi\nThere"),
            ("rtf", rb"{\rtf1\ansi Hello \b bold\b0  world}", "Hello bold world"),
            ("md", "# Title\n\nBody".encode("utf-8-sig"), "# Title\n\nBody"),
        ],
    )
    def test_text_formats(self, file_type, raw, expected):
        from app.services.document_ingestion import extract_pages

        assert extract_pages(raw, file_type)[0][1] == expected


class TestTypeGate:
    def test_every_supported_type_passes(self):
        from app.services.document_ingestion import SUPPORTED_TYPES, unsupported_reason

        assert all(unsupported_reason(t) is None for t in SUPPORTED_TYPES)

    def test_legacy_office_says_what_to_save_as(self):
        from app.services.document_ingestion import unsupported_reason

        assert ".xlsx" in unsupported_reason("xls")
        assert ".docx" in unsupported_reason("doc")

    def test_anything_else_is_refused_with_the_list(self):
        from app.services.document_ingestion import file_type_of, unsupported_reason

        assert file_type_of("archive.tar.gz") == "gz"
        assert file_type_of("README") == ""
        assert "Excel" in unsupported_reason("gz")
