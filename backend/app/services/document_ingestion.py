import asyncio
import csv
import io
import zipfile
import logging
import re
import uuid
from html.parser import HTMLParser

import docx
import pdfplumber
import tiktoken
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.models import DocumentChunk
from app.services.openai_client import embed_texts

logger = logging.getLogger("clardentity.ingestion")

# Section 12.3: ~500-800 tokens per chunk, ~15% overlap.
CHUNK_TOKENS = 650
OVERLAP_TOKENS = 100

_encoding = tiktoken.get_encoding("cl100k_base")

# Everything a workspace upload or a chat attachment may be. The value is the
# content type the file is stored under. Text-like formats all go through
# the same decode; the Office formats each have a reader below.
SUPPORTED_TYPES: dict[str, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "txt": "text/plain",
    "md": "text/markdown",
    "markdown": "text/markdown",
    "csv": "text/csv",
    "tsv": "text/tab-separated-values",
    "json": "application/json",
    "xml": "application/xml",
    "html": "text/html",
    "htm": "text/html",
    "rtf": "application/rtf",
    "log": "text/plain",
    "yaml": "text/plain",
    "yml": "text/plain",
}
# The pre-2007 binary Office formats have no reader here; the message says
# what to do instead of a bare "unsupported".
LEGACY_TYPES: dict[str, str] = {"doc": "docx", "xls": "xlsx", "ppt": "pptx"}


def file_type_of(filename: str) -> str:
    return (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""


def unsupported_reason(file_type: str) -> str | None:
    """None when the type can be read; otherwise the sentence to show."""
    if file_type in SUPPORTED_TYPES:
        return None
    if file_type in LEGACY_TYPES:
        modern = LEGACY_TYPES[file_type]
        return f".{file_type} is the old binary format - save it as .{modern} and attach that."
    return "Unsupported file type. Use PDF, Word, Excel, PowerPoint, or a text file (txt, md, csv, json, html)."


class _TextOnly(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs) -> None:  # noqa: ANN001
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in ("p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def _html_text(raw: str) -> str:
    parser = _TextOnly()
    parser.feed(raw)
    return re.sub(r"\n{3,}", "\n\n", "".join(parser.parts)).strip()


_RTF_CONTROL_RE = re.compile(r"\\[a-z]+-?\d* ?|[{}]|\\'[0-9a-f]{2}")


def _rtf_text(raw: str) -> str:
    return re.sub(r"[ \t]+", " ", _RTF_CONTROL_RE.sub("", raw)).strip()


def _decode(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def _sheet_text(rows) -> str:  # noqa: ANN001
    lines = []
    for row in rows:
        cells = ["" if v is None else str(v).strip() for v in row]
        if any(cells):
            lines.append(" | ".join(cells).rstrip(" |"))
    return "\n".join(lines)


#: Beyond this, table detection is not worth the wall-clock. pdfplumber
#: walks the drawn lines on a page and is an order of magnitude slower than
#: pulling the text layer; a 400-page scan would hold a worker for minutes
#: to find nothing. The text layer still comes out of every page.
_MAX_TABLE_PAGES = 60


def pdf_page_text(page, plumber_page=None) -> str:  # noqa: ANN001
    """One PDF page as text, with any tables laid out as rows.

    A PDF has no idea it contains a table. The text layer is a bag of
    positioned strings, so pulling it in reading order turns a price list
    into a run of numbers with no way to tell which belongs to which drink -
    which is exactly how Kerala alcohol prices came back, Bevco publishing
    them as PDF tables and this reading them linearly.

    So any table found on the page is appended as ' | '-separated rows,
    the same shape a spreadsheet gets, where a figure keeps its column
    neighbours. Appended rather than substituted: table detection misses
    tables ruled only by whitespace, and the prose around them is often
    where the units and the date live.
    """
    text = page.extract_text() or ""
    if plumber_page is None:
        return text

    rendered = []
    try:
        for table in plumber_page.extract_tables() or []:
            laid_out = _sheet_text(table)
            # A "table" of one cell is a text box pdfplumber mistook for
            # one, and repeating it under the prose helps nobody.
            if laid_out and laid_out.count("\n") >= 1:
                rendered.append(laid_out)
    except Exception:  # noqa: BLE001 - a page that will not parse is still a page
        logger.warning("could not read tables from a pdf page", exc_info=True)

    if not rendered:
        return text
    return "\n\n".join([text, *rendered]).strip()


def _pdf_pages(file_bytes: bytes) -> list[tuple[int | None, str]]:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = list(reader.pages)

    plumber_pages: list = [None] * len(pages)
    if len(pages) <= _MAX_TABLE_PAGES:
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                # Zipped by position, and only as far as both agree: the two
                # readers can disagree about page count on a damaged file,
                # and pairing page 3's text with page 4's table would be
                # worse than having no tables at all.
                for i, plumber_page in enumerate(pdf.pages[: len(pages)]):
                    plumber_pages[i] = plumber_page
                return [
                    (i + 1, pdf_page_text(page, plumber_pages[i]))
                    for i, page in enumerate(pages)
                ]
        except Exception:  # noqa: BLE001 - fall back to the text layer alone
            logger.warning("table-aware pdf read failed; using the text layer", exc_info=True)

    return [(i + 1, pdf_page_text(page)) for i, page in enumerate(pages)]


#: An Office file is a zip; a small one can expand to gigabytes of XML (a
#: "zip bomb"). Checked on the zip's own directory before any parser opens it.
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024
MAX_ZIP_ENTRIES = 10_000
#: Text kept from one file before chunking and embedding. Past this the
#: embedding cost grows with no benefit to the answer.
MAX_TEXT_CHARS = 2_000_000


def _check_zip(file_bytes: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > MAX_UNZIPPED_BYTES:
            raise ValueError("file expands to more than the allowed size")


def extract_pages(file_bytes: bytes, file_type: str) -> list[tuple[int | None, str]]:
    """Returns (page_number, text) pairs. page_number is 1-indexed for PDFs
    and slide decks (so citations can point at a real page); the rest have
    no page concept. A spreadsheet yields one entry per sheet, rows laid
    out as ' | '-separated cells so a figure keeps its column neighbours.
    """
    if file_type == "pdf":
        return _pdf_pages(file_bytes)

    if file_type in ("docx", "xlsx", "xlsm", "pptx"):
        _check_zip(file_bytes)

    if file_type == "docx":
        document = docx.Document(io.BytesIO(file_bytes))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            parts.append(_sheet_text([c.text for c in row.cells] for row in table.rows))
        return [(None, "\n\n".join(p for p in parts if p.strip()))]

    if file_type in ("xlsx", "xlsm"):
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        pages = []
        for sheet in workbook.worksheets:
            text = _sheet_text(sheet.iter_rows(values_only=True))
            if text:
                pages.append((None, f"Sheet: {sheet.title}\n{text}"))
        return pages

    if file_type == "pptx":
        deck = Presentation(io.BytesIO(file_bytes))
        pages = []
        for number, slide in enumerate(deck.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    texts.append(_sheet_text([c.text for c in row.cells] for row in shape.table.rows))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                texts.append(slide.notes_slide.notes_text_frame.text)
            pages.append((number, "\n".join(t for t in texts if t.strip())))
        return pages

    raw = _decode(file_bytes)
    if file_type in ("html", "htm"):
        return [(None, _html_text(raw))]
    if file_type == "rtf":
        return [(None, _rtf_text(raw))]
    if file_type in ("csv", "tsv"):
        dialect = "excel-tab" if file_type == "tsv" else "excel"
        return [(None, _sheet_text(csv.reader(io.StringIO(raw), dialect=dialect)))]
    return [(None, raw)]


def chunk_text(text: str) -> list[str]:
    tokens = _encoding.encode(text)
    if not tokens:
        return []

    chunks: list[str] = []
    step = CHUNK_TOKENS - OVERLAP_TOKENS
    start = 0
    while start < len(tokens):
        chunk_tokens = tokens[start : start + CHUNK_TOKENS]
        chunks.append(_encoding.decode(chunk_tokens))
        if start + CHUNK_TOKENS >= len(tokens):
            break
        start += step
    return chunks


async def build_chunks(document_id: uuid.UUID, file_bytes: bytes, file_type: str) -> list[DocumentChunk]:
    """Extract, chunk and embed one file: the rows to insert, not yet added
    to any session. Empty when the file had no readable text. Shared by the
    upload worker and the chat attachment path so both read the same
    formats the same way.
    """
    # Parsing is synchronous and can take seconds on a large file: off the
    # event loop, so one upload does not stall every other request.
    pages = await asyncio.to_thread(extract_pages, file_bytes, file_type)
    kept: list[tuple[int | None, str]] = []
    budget = MAX_TEXT_CHARS
    for page_number, page_text in pages:
        if budget <= 0:
            break
        kept.append((page_number, page_text[:budget]))
        budget -= len(page_text)
    pages = kept
    entries: list[tuple[int | None, str]] = [
        (page_number, chunk)
        for page_number, page_text in pages
        for chunk in chunk_text(page_text)
        if chunk.strip()
    ]
    if not entries:
        return []
    embeddings = await embed_texts([content for _, content in entries])
    return [
        DocumentChunk(
            document_id=document_id,
            chunk_index=index,
            content=content,
            embedding=embedding,
            page_number=page_number,
        )
        for index, ((page_number, content), embedding) in enumerate(zip(entries, embeddings))
    ]
