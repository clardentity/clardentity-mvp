"""Turning a Creative-mode answer into an actual file.

The chat model already wrote the content; this asks a second, structured call
to re-shape that same content into whatever a Word document, a slide deck, or
a spreadsheet actually looks like - a title plus paragraphs, a title plus
slides of bullets, or a table - then renders that structure with python-docx
/ python-pptx / openpyxl. Generated fresh per request and streamed straight
back in the response (see api/chat.py's export_file), the same "no persisted
copy" choice audio.py's TTS already makes - a second export costs one more
call, not a stored file to manage.
"""

import io

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from app.services.anthropic_client import generate_structured
from app.services.output_cleanup import clean_output

EXPORT_FORMATS = ("docx", "pptx", "xlsx")

_FORMAT_LABEL = {
    "docx": "a Word document: a title plus a sequence of sections, each with "
    "an optional heading and a body paragraph. Fill `sections`; leave "
    "`slides` and `table` as empty.",
    "pptx": "a slide deck: a title plus a sequence of slides, each with a "
    "heading and a short list of bullet points. Fill `slides`; leave "
    "`sections` and `table` as empty.",
    "xlsx": "a spreadsheet: a title plus one table with column headers and "
    "rows of string values. Fill `table`; leave `sections` and `slides` "
    "as empty.",
}

_INSTRUCTIONS = (
    "Reshape the given text into the structure a {kind} actually needs - "
    "{detail}\n\n"
    "Keep the writing itself as close to the original as sense allows: this "
    "is a change of shape, not a rewrite. Give the whole thing a short, "
    "specific title - never a generic one like 'Document' or 'Summary'."
)

_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string", "description": "Short, specific title for the file."},
        "sections": {
            "type": "array",
            "description": "Word document only - a title plus paragraphs.",
            "items": {
                "type": "object",
                "properties": {
                    "heading": {
                        # anyOf, not a nullable type array: Anthropic's
                        # structured-output validator rejects the latter (see
                        # anthropic_client._portable_schema, which only
                        # rewrites nullable *enums* to this form - a bare
                        # nullable type without an enum falls through
                        # unconverted, so it's written this way from the
                        # start rather than relying on that rewrite).
                        "anyOf": [{"type": "string"}, {"type": "null"}],
                        "description": "Section heading, or null for an untitled paragraph.",
                    },
                    "body": {"type": "string"},
                },
                "required": ["heading", "body"],
                "additionalProperties": False,
            },
        },
        "slides": {
            "type": "array",
            "description": "Slide deck only - one entry per slide after the title slide.",
            "items": {
                "type": "object",
                "properties": {
                    "heading": {"type": "string"},
                    "bullets": {"type": "array", "items": {"type": "string"}},
                },
                "required": ["heading", "bullets"],
                "additionalProperties": False,
            },
        },
        "table": {
            "type": "object",
            "description": "Spreadsheet only - one table with headers and rows.",
            "properties": {
                "headers": {"type": "array", "items": {"type": "string"}},
                "rows": {
                    "type": "array",
                    "items": {"type": "array", "items": {"type": "string"}},
                },
            },
            "required": ["headers", "rows"],
            "additionalProperties": False,
        },
    },
    "required": ["title", "sections", "slides", "table"],
    "additionalProperties": False,
}


async def build_outline(content: str, export_format: str) -> dict:
    kind = {"docx": "Word document", "pptx": "slide deck", "xlsx": "spreadsheet"}[export_format]
    instructions = _INSTRUCTIONS.format(kind=kind, detail=_FORMAT_LABEL[export_format])
    return await generate_structured(
        instructions=instructions,
        input_text=content,
        schema=_SCHEMA,
        schema_name="office_export_outline",
        fast=True,
    )


def _safe_filename(title: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch in (" ", "-", "_") else "" for ch in title).strip()
    return cleaned[:80] or "document"


def render_docx(outline: dict) -> bytes:
    doc = Document()
    doc.add_heading(clean_output(outline.get("title") or "Document"), level=0)
    for section in outline.get("sections") or []:
        heading = section.get("heading")
        if heading:
            doc.add_heading(clean_output(heading), level=1)
        body = section.get("body")
        if body:
            doc.add_paragraph(clean_output(body))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pptx(outline: dict) -> bytes:
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = clean_output(outline.get("title") or "Presentation")

    bullet_layout = prs.slide_layouts[1]
    for slide_data in outline.get("slides") or []:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = clean_output(slide_data.get("heading") or "")
        bullets = [clean_output(b) for b in (slide_data.get("bullets") or []) if b]
        if not bullets:
            continue
        text_frame = slide.placeholders[1].text_frame
        text_frame.text = bullets[0]
        for bullet in bullets[1:]:
            paragraph = text_frame.add_paragraph()
            paragraph.text = bullet
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def render_xlsx(outline: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _safe_filename(outline.get("title") or "Sheet1")[:31] or "Sheet1"
    table = outline.get("table") or {}
    headers = table.get("headers") or []
    if headers:
        sheet.append([clean_output(h) for h in headers])
    for row in table.get("rows") or []:
        sheet.append([clean_output(cell) for cell in row])
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


_RENDERERS = {"docx": render_docx, "pptx": render_pptx, "xlsx": render_xlsx}

MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


async def export_file(content: str, export_format: str) -> tuple[bytes, str]:
    """Returns (file_bytes, filename). Exceptions propagate as-is - the
    caller (api/chat.py) needs the original exception, not a wrapped one, to
    tell a provider outage apart from a genuine bug via
    is_provider_unavailable_error."""
    outline = await build_outline(content, export_format)
    file_bytes = _RENDERERS[export_format](outline)
    filename = f"{_safe_filename(outline.get('title') or 'document')}.{export_format}"
    return file_bytes, filename
